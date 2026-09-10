#!/usr/bin/env python3
"""Safe deterministic XLSX handoff for Forge's versioned component tables.

The workbook is an editor, not the source of truth.  It embeds only the exact
Git source reference and table contract needed to return it.  The gateway
materializes that ref and rebuilds the trusted three-way-merge baseline.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import math
import os
from pathlib import Path, PurePosixPath
import re
import tempfile
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


FORMAT = "forge-tabular-workbook"
VERSION = 1
MAX_INPUT_BYTES = 10 * 1024 * 1024
MAX_ZIP_ENTRIES = 2048
MAX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
MAX_SHEETS = 32
MAX_ROWS = 5001
MAX_COLS = 256
MAX_CELLS = 250_000
MAX_CELL_CHARS = 100_000
META_SHEET = "_forge_manifest"
META_MARKER = "FORGE WORKBOOK METADATA — DO NOT EDIT"
FIXED_TIME = dt.datetime(2000, 1, 1, 0, 0, 0)
ZIP_TIME = (2000, 1, 1, 0, 0, 0)


class WorkbookError(ValueError):
    pass


def fail(message: str) -> None:
    raise WorkbookError(message)


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def safe_zip(path: Path) -> None:
    if not path.is_file() or path.stat().st_size > MAX_INPUT_BYTES:
        fail("workbook is missing or larger than 10 MB")
    if not zipfile.is_zipfile(path):
        fail("file is not an OOXML .xlsx workbook")
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > MAX_ZIP_ENTRIES:
            fail("workbook has too many package entries")
        total = 0
        for info in infos:
            name = info.filename
            pure = PurePosixPath(name)
            if name.startswith(("/", "\\")) or "\\" in name or ".." in pure.parts:
                fail("workbook contains an unsafe package path")
            total += info.file_size
            if total > MAX_UNCOMPRESSED_BYTES:
                fail("workbook expands beyond the 64 MB safety limit")
            if info.file_size > 1024 * 1024 and info.compress_size and info.file_size / info.compress_size > 200:
                fail("workbook contains a suspiciously compressed package entry")
            lowered = name.lower()
            if lowered.endswith("vbaproject.bin") or lowered.startswith("xl/externallinks/"):
                fail("macros and external workbook links are not accepted")


def canonical_zip(source: Path, target: Path) -> None:
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(
        target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as outgoing:
        for name in sorted(incoming.namelist()):
            info = zipfile.ZipInfo(name, ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o100644 << 16
            data = incoming.read(name)
            if name == "docProps/core.xml":
                data = re.sub(
                    br"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    br"\g<1>2000-01-01T00:00:00Z\g<2>", data,
                )
            outgoing.writestr(info, data, compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)


def chunks(value: str, size: int = 30_000) -> list[str]:
    return [value[index:index + size] for index in range(0, len(value), size)] or [""]


def csv_rows(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text, newline="")))
    if not rows or "id" not in rows[0]:
        fail("editable table is missing its stable id column")
    return rows


def workbook_metadata(project_zip: Path) -> tuple[dict, dict[str, list[list[str]]]]:
    with zipfile.ZipFile(project_zip) as archive:
        try:
            project = json.loads(archive.read("manifest.json"))
        except (KeyError, json.JSONDecodeError, UnicodeDecodeError) as error:
            raise WorkbookError("working copy is missing a valid manifest.json") from error
        if project.get("format") != "forge-design-project" or project.get("profile") != "forge-tabular-working-copy":
            fail("input is not a Forge tabular working copy")
        source = project.get("source") or {}
        ref = str(source.get("ref") or "")
        if not ref or not all(char in "0123456789abcdefABCDEF" for char in ref) or not 7 <= len(ref) <= 40:
            fail("working copy needs an exact Git source reference")
        tables: dict[str, list[list[str]]] = {}
        table_meta: dict[str, dict] = {}
        for role in ("cards", "printings", "tokens"):
            spec = (project.get("tables") or {}).get(role)
            if not spec:
                continue
            path = spec.get("editable_path")
            try:
                rows = csv_rows(archive.read(path))
            except KeyError as error:
                raise WorkbookError(f"working copy is missing {path}") from error
            tables[role] = rows
            table_meta[role] = {
                "sheet": role,
                "id": "id",
                "columns": rows[0],
                "csv_sha256": sha256(archive.read(path)),
            }
        if "cards" not in tables or "printings" not in tables:
            fail("working copy needs cards and printings tables")
        metadata = {
            "format": FORMAT,
            "version": VERSION,
            "profile": "forge-tabular-working-copy",
            "game": project.get("game") or {},
            "source": {"ref": ref, "hash": source.get("hash")},
            "tables": table_meta,
            "rules": {
                "stable_id_required": True,
                "formulas_allowed": False,
                "atomic_return": True,
            },
        }
        return metadata, tables


def write_table_sheet(workbook: Workbook, role: str, rows: list[list[str]]) -> None:
    sheet = workbook.create_sheet(role)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(rows[0]))}{max(1, len(rows))}"
    header_fill = PatternFill("solid", fgColor="182336")
    for row_index, values in enumerate(rows, start=1):
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            # Prevent a canonical string beginning with =, +, -, or @ from
            # becoming executable spreadsheet content on export.
            cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=row_index > 1)
            if row_index == 1:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="center")
            else:
                cell.number_format = "@"
    sheet.row_dimensions[1].height = 24
    for column_index, header in enumerate(rows[0], start=1):
        observed = [str(row[column_index - 1]) for row in rows[:101] if len(row) >= column_index]
        width = max([len(str(header)), *[min(48, max((len(line) for line in value.splitlines()), default=0)) for value in observed]])
        sheet.column_dimensions[get_column_letter(column_index)].width = min(48, max(10, width + 2))


def export_workbook(project_zip: Path, output: Path) -> dict:
    metadata, tables = workbook_metadata(project_zip)
    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    instructions = [
        ["Forge versioned workbook"],
        ["Edit the visible table sheets, then return this same .xlsx file through Forge Design."],
        ["Keep every stable id. IDs connect renames to history, printings, decks, diffs, and credit."],
        ["Cards, printings, and tokens return together as one dry run and one commit or pull request."],
        ["Do not add formulas, macros, or external workbook links. Forge rejects them instead of guessing values."],
        ["The hidden Forge sheet stores the exact Git ref. Forge rebuilds the trusted baseline from Git on return."],
        [f"Game: {metadata['game'].get('title') or metadata['game'].get('id') or ''}"],
        [f"Source ref: {metadata['source']['ref']}"],
    ]
    for row in instructions:
        readme.append(row)
    readme.column_dimensions["A"].width = 110
    readme["A1"].font = Font(size=18, bold=True, color="182336")
    for cell in readme["A"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for role, rows in tables.items():
        write_table_sheet(workbook, role, rows)
    meta = workbook.create_sheet(META_SHEET)
    meta.sheet_state = "veryHidden"
    meta["A1"] = META_MARKER
    payload = json.dumps(metadata, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    for index, part in enumerate(chunks(payload), start=2):
        meta.cell(row=index, column=1, value=part)
    workbook.properties.creator = "Forge"
    workbook.properties.lastModifiedBy = "Forge"
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME
    workbook.properties.title = f"{metadata['game'].get('title') or metadata['game'].get('id') or 'Forge'} data working copy"
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="forge-workbook-") as temp:
        raw = Path(temp) / "raw.xlsx"
        workbook.save(raw)
        canonical_zip(raw, output)
    return metadata


def value_text(value, data_type: str, role: str, coordinate: str) -> str:
    if data_type == "f":
        fail(f"{role} {coordinate} contains a formula; paste its literal value before returning")
    if data_type == "e":
        fail(f"{role} {coordinate} contains a spreadsheet error")
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            fail(f"{role} {coordinate} contains a non-finite number")
        return str(int(value)) if value.is_integer() else format(value, ".15g")
    text = str(value).replace("\x00", "")
    if len(text) > MAX_CELL_CHARS:
        fail(f"{role} {coordinate} is larger than the 100,000 character cell limit")
    return text


def read_metadata(workbook) -> dict:
    if META_SHEET not in workbook.sheetnames:
        fail("workbook has lost its hidden Forge baseline marker; download a fresh working copy")
    sheet = workbook[META_SHEET]
    if sheet.max_column != 1 or sheet.max_row > 1000:
        fail("workbook Forge metadata exceeds its bounded single-column envelope")
    if sheet["A1"].value != META_MARKER:
        fail("workbook Forge metadata marker is invalid")
    payload = "".join(str(sheet.cell(row=row, column=1).value or "") for row in range(2, sheet.max_row + 1))
    try:
        metadata = json.loads(payload)
    except json.JSONDecodeError as error:
        raise WorkbookError("workbook Forge metadata is not valid JSON") from error
    if metadata.get("format") != FORMAT or metadata.get("version") != VERSION:
        fail(f"unsupported Forge workbook format {metadata.get('format')} v{metadata.get('version')}")
    ref = str((metadata.get("source") or {}).get("ref") or "")
    if not ref or not all(char in "0123456789abcdefABCDEF" for char in ref) or not 7 <= len(ref) <= 40:
        fail("workbook has no valid exact Git source reference")
    return metadata


def import_workbook(path: Path) -> dict:
    safe_zip(path)
    try:
        workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    except Exception as error:
        raise WorkbookError(f"could not open workbook: {error}") from error
    if len(workbook.sheetnames) > MAX_SHEETS:
        fail("workbook has more than 32 worksheets")
    metadata = read_metadata(workbook)
    tables: dict[str, dict] = {}
    total_cells = 0
    for role, spec in (metadata.get("tables") or {}).items():
        if role not in ("cards", "printings", "tokens"):
            fail(f"workbook declares unsupported table role '{role}'")
        name = str(spec.get("sheet") or "")
        if name not in workbook.sheetnames or name == META_SHEET:
            fail(f"workbook is missing its '{name or role}' {role} sheet")
        sheet = workbook[name]
        if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLS:
            fail(f"{role} exceeds the {MAX_ROWS - 1} row or {MAX_COLS} column limit")
        total_cells += sheet.max_row * sheet.max_column
        if total_cells > MAX_CELLS:
            fail("workbook tables exceed the 250,000 cell safety limit")
        rows: list[list[str]] = []
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=sheet.max_column):
            values = [value_text(cell.value, cell.data_type, role, cell.coordinate) for cell in row]
            while values and values[-1] == "":
                values.pop()
            rows.append(values)
        while rows and not any(rows[-1]):
            rows.pop()
        if not rows:
            fail(f"{role} sheet has no header row")
        headers = [value.strip() for value in rows[0]]
        if "id" not in headers:
            fail(f"{role} sheet needs the stable id column")
        if len(headers) != len(set(headers)):
            fail(f"{role} sheet has duplicate columns")
        stream = io.StringIO(newline="")
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(headers)
        for values in rows[1:]:
            if any(value != "" for value in values):
                writer.writerow(values + [""] * (len(headers) - len(values)))
        tables[role] = {"sheet": name, "csv": stream.getvalue(), "rows": max(0, len(rows) - 1), "columns": headers}
    if "cards" not in tables or "printings" not in tables:
        fail("workbook needs cards and printings tables")
    return {
        "format": FORMAT,
        "version": VERSION,
        "profile": metadata.get("profile"),
        "game": metadata.get("game") or {},
        "source": metadata.get("source") or {},
        "tables": tables,
        "warnings": ["Cell comments, styling, and unrelated worksheets do not enter Forge; canonical table values do."],
    }


def inspect_candidate(path: Path) -> dict:
    """Read ordinary user-owned XLSX tabs for the existing new-game mapper."""
    safe_zip(path)
    try:
        workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    except Exception as error:
        raise WorkbookError(f"could not open workbook: {error}") from error
    if len(workbook.sheetnames) > MAX_SHEETS:
        fail("workbook has more than 32 worksheets")
    sheets: list[dict] = []
    total_cells = 0
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible" or sheet.title == META_SHEET:
            continue
        if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLS:
            fail(f"{sheet.title} exceeds the {MAX_ROWS - 1} row or {MAX_COLS} column limit")
        total_cells += sheet.max_row * sheet.max_column
        if total_cells > MAX_CELLS:
            fail("workbook tables exceed the 250,000 cell safety limit")
        rows: list[list[str]] = []
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=sheet.max_column):
            values = [value_text(cell.value, cell.data_type, sheet.title, cell.coordinate) for cell in row]
            while values and values[-1] == "":
                values.pop()
            rows.append(values)
        while rows and not any(rows[-1]):
            rows.pop()
        if not rows or not any(value.strip() for value in rows[0]):
            continue
        headers = [value.strip() for value in rows[0]]
        if len(headers) != len(set(headers)):
            fail(f"{sheet.title} has duplicate columns")
        stream = io.StringIO(newline="")
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(headers)
        data_rows = 0
        for values in rows[1:]:
            if any(value != "" for value in values):
                writer.writerow(values + [""] * (len(headers) - len(values)))
                data_rows += 1
        normalized = {header.strip().lower().replace(" ", "_") for header in headers}
        score = (8 if "name" in normalized or "card_name" in normalized or "title" in normalized else 0) \
            + (4 if "id" in normalized else 0) + (2 if "type" in normalized else 0) \
            + (2 if "text" in normalized or "rules_text" in normalized else 0) + min(data_rows, 5)
        sheets.append({"name": sheet.title, "csv": stream.getvalue(), "rows": data_rows, "columns": headers, "card_score": score})
    if not sheets:
        fail("workbook has no visible non-empty table sheets")
    suggested = max(sheets, key=lambda sheet: (sheet["card_score"], sheet["rows"]))["name"]
    return {
        "format": "forge-workbook-candidate",
        "version": 1,
        "sheets": sheets,
        "suggested_sheet": suggested,
        "warnings": ["Choose the card-data tab. Other workbook tabs remain untouched and can be brought into the traced Forge workbook after project creation."],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    export = sub.add_parser("export")
    export.add_argument("project_zip", type=Path)
    export.add_argument("output", type=Path)
    inspect = sub.add_parser("inspect")
    inspect.add_argument("workbook", type=Path)
    inspect.add_argument("output", nargs="?", type=Path)
    candidate = sub.add_parser("inspect-candidate")
    candidate.add_argument("workbook", type=Path)
    candidate.add_argument("output", nargs="?", type=Path)
    args = parser.parse_args()
    try:
        if args.command == "export":
            result = export_workbook(args.project_zip, args.output)
        elif args.command == "inspect":
            result = import_workbook(args.workbook)
        else:
            result = inspect_candidate(args.workbook)
        encoded = json.dumps(result, sort_keys=True, ensure_ascii=False, separators=(",", ":")) + "\n"
        if args.command in ("inspect", "inspect-candidate") and args.output:
            args.output.write_text(encoded, encoding="utf-8")
        else:
            print(encoded, end="")
        return 0
    except (WorkbookError, OSError, zipfile.BadZipFile) as error:
        print(json.dumps({"error": str(error)}, ensure_ascii=False), file=os.sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
